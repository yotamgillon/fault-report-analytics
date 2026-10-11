"""Build the Excel report with openpyxl: real formulas over a clean Data table.

How the workbook fits together
- ``Data`` is an Excel Table (tblFaults): one row per ticket. Its helper columns (resolved,
  days to resolve, age, SLA, month) are formulas, and vendor tier / SLA are joined from the
  ``Vendors`` table with XLOOKUP.
- ``By_Vendor``, ``By_System``, ``By_Site``, ``Monthly`` and ``Backlog`` are pivot-style summaries
  made of COUNTIFS / AVERAGEIFS and array-formula medians over tblFaults, so they recalculate when
  the data or the Settings change.
- ``Dashboard`` has two dropdown filters (vendor, site), KPI cards driven by them, and charts.
- Only labels, the sort order of the rows and pipeline outputs that need sequential logic (the
  repeat-fault flags, computed by ``metrics.repeat_faults``) are values; every number a reader
  sees on a summary sheet is a formula.

Two compatibility choices worth knowing
- Conditional medians use an array formula, MEDIAN(IF(...)). AGGREGATE(16,6,x/(cond),0.5) is
  valid in Excel but returns #VALUE! in LibreOffice, which is what CI recalculates with.
- XLOOKUP is written as _xlfn.XLOOKUP (the prefix Excel stores) and needs LibreOffice >= 24.8 to
  be recalculated outside Excel.
"""

import json
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableStyleInfo

from field_ops import config
from field_ops.metrics import AGE_BUCKETS, MIN_SAMPLE, Metrics

FONT = "Arial"
INK, INK2, MUTED = "0B0B0B", "52514E", "898781"
HEADER_FILL = "1F2937"
INPUT_FILL, INPUT_INK = "FFF8C5", "0000FF"  # yellow fill + blue text = a cell you may edit
BLUE, ORANGE = "2A78D6", "EB6834"  # validated categorical slots 1 and 2
RED_FILL, RED_INK = "FBE0E0", "A12121"
AMBER_FILL, AMBER_INK = "FFF1CC", "7A5200"
GRID = "E1E0D9"
CARD_FILL, CARD_LINE = "FCFCFB", "C3C2B7"

FMT_DATE, FMT_1DP, FMT_INT = "yyyy-mm-dd", "0.0", "#,##0"
FMT_PCT, FMT_X = "0.0%", '0.00"x"'
FMT_PTS = '+0.0" pts";-0.0" pts";0.0" pts"'
FMT_MONTH = "mmm yyyy"

ALL_VENDORS, ALL_SITES = "All vendors", "All sites"
BUCKETS = [label for label, _, _ in AGE_BUCKETS]
DATA_TABLE, VENDOR_TABLE = "tblFaults", "tblVendors"


def f(**kw) -> Font:
    kw.setdefault("name", FONT)
    kw.setdefault("size", 10)
    return Font(**kw)


def fill(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


def col(name: str) -> str:
    """Structured reference to a column of the Data table."""
    return f"{DATA_TABLE}[{name}]"


# ---------------------------------------------------------------------------------------------
# Layout record: where everything ended up (used by tests to read values back)
# ---------------------------------------------------------------------------------------------
@dataclass
class Layout:
    path: Path
    n_rows: int = 0
    cols: dict[str, dict[str, str]] = field(default_factory=dict)  # sheet -> key -> letter
    rows: dict[str, dict] = field(default_factory=dict)  # sheet -> {"index": {key: row}, ...}
    kpi: dict[str, str] = field(default_factory=dict)  # KPI name -> Dashboard cell
    selectors: dict[str, str] = field(default_factory=dict)


@dataclass
class Column:
    key: str
    header: str
    width: float
    fmt: str | None = None
    formula: Callable[[int, dict[str, str]], str] | None = None
    array: bool = False


def _put(ws, ref: str, value, *, array: bool = False, fmt: str | None = None, **style) -> None:
    ws[ref] = ArrayFormula(ref, value) if array else value
    cell = ws[ref]
    cell.font = style.pop("font", f())
    if fmt:
        cell.number_format = fmt
    for attr, val in style.items():
        setattr(cell, attr, val)


def _header_row(ws, row: int, headers: list[str], first_col: int = 1) -> None:
    for offset, text in enumerate(headers):
        cell = ws.cell(row=row, column=first_col + offset, value=text)
        cell.font = f(bold=True, color="FFFFFF")
        cell.fill = fill(HEADER_FILL)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = 42


def _title(ws, text: str, note: str) -> None:
    ws["A1"] = text
    ws["A1"].font = f(size=14, bold=True, color=INK)
    ws["A2"] = note
    ws["A2"].font = f(color=INK2)
    ws["A2"].alignment = Alignment(wrap_text=False)
    ws.sheet_view.showGridLines = False


def _widths(ws, widths: dict[str, float]) -> None:
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width


def _name(wb: Workbook, name: str, ref: str) -> None:
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


# ---------------------------------------------------------------------------------------------
# Settings, Lists, Vendors
# ---------------------------------------------------------------------------------------------
def _build_settings(wb: Workbook, m: Metrics) -> None:
    ws = wb.create_sheet("Settings")
    _title(ws, "Settings and notes", "Yellow cells with blue text are the only inputs.")
    inputs = [
        ("Reporting date (AsOfDate)", m.as_of, FMT_DATE, "AsOfDate", True,
         "Fixed, not today's date, so ages and the backlog are reproducible."),
        ("Minimum sample (MinSample)", MIN_SAMPLE, "0", "MinSample", True,
         "Groups with fewer resolved tickets are marked 'low sample' and should not be ranked."),
        ("SLA breach alert (BreachAlert)", 0.4, "0%", "BreachAlert", True,
         "Breach rates above this are highlighted red. Arbitrary demo threshold."),
        ("Repeat-rate alert (RepeatAlert)", 1.25, '0.00"x"', "RepeatAlert", True,
         "Systems repeating at or above this multiple of the fleet rate are highlighted amber."),
        ("Repeat window in days (RepeatWindow)", m.repeat.summary["window_days"], "0",
         "RepeatWindow", False,
         "Documentation only: the repeat flags are computed by the pipeline "
         "(field-ops metrics --repeat-window N), so editing this cell changes nothing."),
    ]  # fmt: skip
    for i, (label, value, fmt, name, editable, note) in enumerate(inputs):
        row = 4 + i
        ws.cell(row=row, column=2, value=label).font = f(bold=True)
        cell = ws.cell(row=row, column=3, value=value)
        cell.number_format = fmt
        cell.font = f(color=INPUT_INK if editable else INK)
        if editable:
            cell.fill = fill(INPUT_FILL)
        ws.cell(row=row, column=4, value=note).font = f(color=INK2)
        _name(wb, name, f"Settings!$C${row}")
    notes = [
        "How to read this workbook",
        "- Dashboard: pick a vendor and site; the KPI cards recalculate. Charts show the fleet.",
        "- Data: one row per ticket (an Excel Table). Grey-header columns are formulas.",
        "- Vendors: reference table (tier, SLA target, region, contract). Joined with XLOOKUP.",
        "- By_Vendor, By_System, By_Site, Monthly, Backlog: formula summaries over the Data table.",
        "- Data_Quality: what the cleaning step fixed (pipeline values, an audit log).",
        "",
        "Definitions",
        "- Resolved: status Closed with a usable close date (not flagged closed_no_close_date, "
        "close_before_open or date_invalid). Only resolved tickets enter resolution-time numbers.",
        "- Open: status Open or In Progress. Age = reporting date minus open date.",
        "- Low sample: fewer than MinSample resolved tickets. "
        "Medians come first; means sit beside them.",
        "- SLA breach: a resolved ticket that took longer than its vendor's sla_target_days.",
        "- Repeat fault: same system and category opening on or after the previous fault's close "
        "date and within the repeat window; a fault opened while the previous one was still open "
        "is not a repeat. Unclassified or text-mismatch categories are excluded.",
        "",
        "Assumptions",
        "- All data is synthetic. Vendor tiers, SLA targets, regions and contract types are "
        "arbitrary demo values (data/reference/vendors.csv).",
        "- Repeat flags (Data columns repeat_counted / is_repeat) come from the Python pipeline "
        "because the rule needs the previous ticket per system and category; the rates on the "
        "summary sheets are formulas over those flags.",
        "- Summary rows are sorted by the pipeline at build time; the numbers themselves are live.",
    ]
    for i, text in enumerate(notes):
        cell = ws.cell(row=11 + i, column=2, value=text)
        cell.font = f(bold=text in ("How to read this workbook", "Definitions", "Assumptions"))
    _widths(ws, {"A": 2, "B": 38, "C": 14, "D": 100})
    ws["C4"].alignment = Alignment(horizontal="right")
    ws.sheet_properties.tabColor = MUTED


def _build_lists(wb: Workbook, m: Metrics) -> None:
    ws = wb.create_sheet("Lists")
    vendors = [ALL_VENDORS, *sorted(m.prepared["vendor"].unique())]
    sites = [ALL_SITES, *sorted(m.prepared["site"].unique())]
    lists = {
        "lstVendors": vendors,
        "lstSites": sites,
        "lstStatus": list(config.STATUSES),
        "lstSeverity": list(config.SEVERITIES),
        "lstCategory": [*config.CATEGORIES, config.UNCLASSIFIED],
    }
    for i, (name, values) in enumerate(lists.items(), start=1):
        letter = get_column_letter(i)
        ws.cell(row=1, column=i, value=name).font = f(bold=True)
        for r, value in enumerate(values, start=2):
            ws.cell(row=r, column=i, value=value).font = f()
        _name(wb, name, f"Lists!${letter}$2:${letter}${len(values) + 1}")
    ws.sheet_state = "hidden"


def _build_vendors(wb: Workbook, vendors: pd.DataFrame) -> None:
    ws = wb.create_sheet("Vendors")
    _title(
        ws,
        "Vendor reference",
        "Lookup table joined onto every ticket with XLOOKUP. Synthetic demo values: edit freely.",
    )
    headers = ["vendor", "tier", "sla_target_days", "support_region", "contract_type"]
    _header_row(ws, 4, headers)
    for r, row in enumerate(vendors[headers].itertuples(index=False), start=5):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = f(color=INPUT_INK)
    last = 4 + len(vendors)
    ws.add_table(
        Table(displayName=VENDOR_TABLE, ref=f"A4:E{last}", tableStyleInfo=TableStyleInfo(
            name="TableStyleLight1", showRowStripes=True))
    )  # fmt: skip
    _widths(ws, {"A": 14, "B": 10, "C": 16, "D": 18, "E": 18})
    ws.freeze_panes = "A5"
    ws.sheet_properties.tabColor = MUTED


# ---------------------------------------------------------------------------------------------
# Data table
# ---------------------------------------------------------------------------------------------
VALUE_COLUMNS = [  # (key, width, number format)
    ("report_id", 11, None), ("date_opened", 12, FMT_DATE), ("date_closed", 12, FMT_DATE),
    ("site", 13, None), ("system_type", 22, None), ("system_id", 10, None),
    ("vendor", 10, None), ("fault_category", 15, None), ("severity", 9, None),
    ("status", 11, None), ("category_source", 16, None), ("text_category", 15, None),
    ("dq_flags", 30, None), ("description", 52, None), ("reported_by", 14, None),
    ("dup_count", 9, "0"),
]  # fmt: skip
FORMULA_COLUMNS = [
    ("resolved", 9, "0"), ("days_to_resolve", 11, "0"), ("is_open", 8, "0"),
    ("age_days", 9, "0"), ("age_bucket", 12, None), ("vendor_tier", 10, None),
    ("sla_target_days", 11, "0"), ("sla_breach", 10, "0"),
    ("opened_month", 12, FMT_MONTH), ("closed_month", 12, FMT_MONTH),
]  # fmt: skip
PIPELINE_COLUMNS = [
    ("repeat_counted", 11, "0"),
    ("is_repeat", 9, "0"),
    ("repeat_exclusion", 24, None),
]


def _flag_absent(flag: str, flags_cell: str) -> str:
    return f'NOT(ISNUMBER(SEARCH(";{flag};",";"&{flags_cell}&";")))'


def _data_formulas(c: dict[str, str], r: int) -> dict[str, str]:
    cell = {k: f"{v}{r}" for k, v in c.items()}
    no_bad_flags = ",".join(
        _flag_absent(flag, cell["dq_flags"])
        for flag in ("closed_no_close_date", "close_before_open", "date_invalid")
    )
    return {
        "resolved": (
            f'=IF(AND({cell["status"]}="Closed",ISNUMBER({cell["date_opened"]}),'
            f"ISNUMBER({cell['date_closed']}),{cell['date_closed']}>={cell['date_opened']},"
            f"{no_bad_flags}),1,0)"
        ),
        "days_to_resolve": (
            f'=IF({cell["resolved"]}=1,{cell["date_closed"]}-{cell["date_opened"]},"")'
        ),
        "is_open": (
            f'=IF(AND(ISNUMBER({cell["date_opened"]}),OR({cell["status"]}="Open",'
            f'{cell["status"]}="In Progress")),1,0)'
        ),
        "age_days": f'=IF({cell["is_open"]}=1,MAX(0,AsOfDate-{cell["date_opened"]}),"")',
        "age_bucket": (
            f'=IF({cell["age_days"]}="","",IF({cell["age_days"]}<=7,"{BUCKETS[0]}",'
            f'IF({cell["age_days"]}<=30,"{BUCKETS[1]}",IF({cell["age_days"]}<=90,'
            f'"{BUCKETS[2]}","{BUCKETS[3]}"))))'
        ),
        "vendor_tier": (
            f'=_xlfn.XLOOKUP({cell["vendor"]},{VENDOR_TABLE}[vendor],{VENDOR_TABLE}[tier],"Unknown")'
        ),
        "sla_target_days": (
            f"=_xlfn.XLOOKUP({cell['vendor']},{VENDOR_TABLE}[vendor],"
            f'{VENDOR_TABLE}[sla_target_days],"")'
        ),
        "sla_breach": (
            f"=IF(AND({cell['resolved']}=1,ISNUMBER({cell['sla_target_days']})),"
            f'IF({cell["days_to_resolve"]}>{cell["sla_target_days"]},1,0),"")'
        ),
        "opened_month": (
            f"=IF(ISNUMBER({cell['date_opened']}),DATE(YEAR({cell['date_opened']}),"
            f'MONTH({cell["date_opened"]}),1),"")'
        ),
        "closed_month": (
            f"=IF({cell['resolved']}=1,DATE(YEAR({cell['date_closed']}),"
            f'MONTH({cell["date_closed"]}),1),"")'
        ),
    }


def _timestamp(value):
    return None if pd.isna(value) else value.to_pydatetime()


def _build_data(wb: Workbook, m: Metrics, layout: Layout) -> None:
    ws = wb.create_sheet("Data")
    keys = [k for k, _, _ in VALUE_COLUMNS + FORMULA_COLUMNS + PIPELINE_COLUMNS]
    letters = {k: get_column_letter(i) for i, k in enumerate(keys, start=1)}
    layout.cols["Data"] = letters

    frame = m.prepared.join(
        m.repeat.rows.rename(
            columns={"counted": "repeat_counted", "exclusion": "repeat_exclusion"}
        )[["repeat_counted", "is_repeat", "repeat_exclusion"]]
    )
    ws.append(keys)
    for r, row in enumerate(frame.itertuples(index=False), start=2):
        record = dict(zip(frame.columns, row, strict=True))
        formulas = _data_formulas(letters, r)
        values = []
        for key in keys:
            if key in formulas:
                values.append(formulas[key])
            elif key in ("date_opened", "date_closed"):
                values.append(_timestamp(record[key]))
            elif key in ("repeat_counted", "is_repeat"):
                values.append(int(record[key]))
            elif key == "dup_count":
                values.append(int(record[key]))
            else:
                values.append(record[key])
        ws.append(values)
    n = len(frame)
    layout.n_rows = n
    last = n + 1

    all_cols = VALUE_COLUMNS + FORMULA_COLUMNS + PIPELINE_COLUMNS
    for i, (key, width, fmt) in enumerate(all_cols, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = width
        head = ws.cell(row=1, column=i)
        head.font = f(bold=True, color="FFFFFF")
        head.fill = fill(HEADER_FILL if key not in formulas_keys() else "4B5563")
        head.alignment = Alignment(wrap_text=True, vertical="center")
        if fmt:
            for r in range(2, last + 1):
                ws[f"{letter}{r}"].number_format = fmt
    ws.row_dimensions[1].height = 32
    ws.add_table(
        Table(
            displayName=DATA_TABLE,
            ref=f"A1:{get_column_letter(len(keys))}{last}",
            tableStyleInfo=TableStyleInfo(name="TableStyleLight1", showRowStripes=True),
        )
    )
    ws.freeze_panes = "B2"

    for key, name in (("status", "lstStatus"), ("severity", "lstSeverity"),
                      ("fault_category", "lstCategory")):  # fmt: skip
        dv = DataValidation(type="list", formula1=f"={name}", allow_blank=True)
        dv.error, dv.errorTitle = "Pick a value from the list.", "Not allowed"
        ws.add_data_validation(dv)
        dv.add(f"{letters[key]}2:{letters[key]}{last}")

    ws.conditional_formatting.add(
        f"{letters['age_bucket']}2:{letters['age_bucket']}{last}",
        FormulaRule(
            formula=[f'{letters["age_bucket"]}2="{BUCKETS[3]}"'],
            fill=fill(RED_FILL), font=Font(name=FONT, color=RED_INK),
        ),
    )  # fmt: skip
    ws.conditional_formatting.add(
        f"{letters['sla_breach']}2:{letters['sla_breach']}{last}",
        FormulaRule(
            formula=[f'{letters["sla_breach"]}2=1'], fill=fill(RED_FILL),
            font=Font(name=FONT, color=RED_INK),
        ),
    )  # fmt: skip
    ws.conditional_formatting.add(
        f"{letters['dq_flags']}2:{letters['dq_flags']}{last}",
        FormulaRule(
            formula=[f'ISNUMBER(SEARCH("category_text_mismatch",{letters["dq_flags"]}2))'],
            fill=fill(AMBER_FILL), font=Font(name=FONT, color=AMBER_INK),
        ),
    )  # fmt: skip
    ws.sheet_properties.tabColor = "1F2937"


def formulas_keys() -> set[str]:
    return {k for k, _, _ in FORMULA_COLUMNS}


# ---------------------------------------------------------------------------------------------
# Summary sheets
# ---------------------------------------------------------------------------------------------
def _letters(columns: list[Column]) -> dict[str, str]:
    return {c.key: get_column_letter(i) for i, c in enumerate(columns, start=1)}


def _write_summary(
    wb: Workbook,
    name: str,
    title: str,
    note: str,
    columns: list[Column],
    keys: list[str],
    total: dict[str, str] | None,
    layout: Layout,
    first_row: int = 6,
    tab: str = BLUE,
) -> tuple[dict[str, str], dict[str, int], int | None]:
    ws = wb.create_sheet(name)
    _title(ws, title, note)
    letters = _letters(columns)
    _header_row(ws, first_row - 1, [c.header for c in columns])
    index = {key: first_row + i for i, key in enumerate(keys)}
    total_row = first_row + len(keys) if total is not None else None

    for key, r in index.items():
        for c in columns:
            ref = f"{letters[c.key]}{r}"
            if c.formula is None:
                _put(ws, ref, key, font=f(bold=True))
            else:
                _put(ws, ref, c.formula(r, {**letters, "_total": str(total_row)}),
                     array=c.array, fmt=c.fmt)  # fmt: skip
    if total_row is not None:
        for c in columns:
            ref = f"{letters[c.key]}{total_row}"
            expr = total.get(c.key) if total else None
            if c.key == columns[0].key:
                _put(ws, ref, "All", font=f(bold=True), fill=fill("F3F4F6"))
            elif expr is not None:
                _put(ws, ref, expr, fmt=c.fmt, font=f(bold=True), fill=fill("F3F4F6"))
            else:
                ws[ref].fill = fill("F3F4F6")
        ws[f"{letters[columns[0].key]}{total_row}"].border = Border(top=Side(style="thin"))
    for c in columns:
        ws.column_dimensions[letters[c.key]].width = c.width
    ws.freeze_panes = ws[f"B{first_row}"]
    ws.sheet_properties.tabColor = tab
    layout.cols[name] = letters
    layout.rows[name] = {"index": index, "total": total_row, "first": first_row}
    return letters, index, total_row


def _cf_low_sample(ws, letters, first, last, sample_key="sample") -> None:
    rng = f"A{first}:{letters[max(letters, key=lambda k: letters[k])]}{last}"
    ws.conditional_formatting.add(
        rng,
        FormulaRule(
            formula=[f'${letters[sample_key]}{first}="low sample"'],
            font=Font(name=FONT, italic=True, color=MUTED),
        ),
    )


def _cf_scale(ws, ref: str) -> None:
    ws.conditional_formatting.add(
        ref,
        ColorScaleRule(start_type="min", start_color="E8F1FC", end_type="max", end_color="6DA7EC"),
    )


def _cf_bars(ws, ref: str) -> None:
    ws.conditional_formatting.add(
        ref, DataBarRule(start_type="num", start_value=0, end_type="max", color="86B6EF")
    )


def _cf_alert(ws, ref: str, formula: str, kind: str = "red") -> None:
    fill_c, ink = (RED_FILL, RED_INK) if kind == "red" else (AMBER_FILL, AMBER_INK)
    ws.conditional_formatting.add(
        ref, FormulaRule(formula=[formula], fill=fill(fill_c), font=Font(name=FONT, color=ink))
    )


def _ratio(num: str, den: str) -> str:
    """Formula for num/den that stays blank when the denominator is 0."""
    return f'=IF({den}=0,"",{num}/{den})'


def _vendor_columns(first_total: str) -> tuple[list[Column], dict[str, str]]:
    """Columns of By_Vendor plus the formulas of its total row."""

    def lookup(field_name: str):
        return lambda r, c: (
            f'=_xlfn.XLOOKUP(${c["vendor"]}{r},{VENDOR_TABLE}[vendor],{VENDOR_TABLE}[{field_name}],"")'
        )

    def cnt(extra: str):
        return lambda r, c: f"=COUNTIFS({col('vendor')},${c['vendor']}{r}{extra})"

    columns = [
        Column("vendor", "Vendor", 13),
        Column("tier", "Tier", 9, None, lookup("tier")),
        Column("sla_days", "SLA target (days)", 10, "0", lookup("sla_target_days")),
        Column("tickets", "Tickets", 9, FMT_INT, cnt("")),
        Column("resolved", "Resolved (n)", 10, FMT_INT, cnt(f",{col('resolved')},1")),
        Column(
            "median",
            "Median days to resolve",
            12,
            FMT_1DP,
            lambda r, c: (
                f'=IF({c["resolved"]}{r}=0,"",MEDIAN(IF({col("vendor")}=${c["vendor"]}{r},'
                f"{col('days_to_resolve')})))"
            ),
            True,
        ),
        Column(
            "mean",
            "Mean days to resolve",
            12,
            FMT_1DP,
            lambda r, c: (
                f"=IFERROR(AVERAGEIFS({col('days_to_resolve')},{col('vendor')},${c['vendor']}{r},"
                f'{col("resolved")},1),"")'
            ),
        ),
        Column(
            "sample",
            "Sample",
            11,
            None,
            lambda r, c: f'=IF({c["resolved"]}{r}<MinSample,"low sample","")',
        ),
        Column("breaches", "SLA breaches", 10, FMT_INT, cnt(f",{col('sla_breach')},1")),
        Column(
            "breach_rate",
            "SLA breach rate",
            10,
            FMT_PCT,
            lambda r, c: f'=IF({c["resolved"]}{r}=0,"",{c["breaches"]}{r}/{c["resolved"]}{r})',
        ),
        Column(
            "rep_counted",
            "Faults counted for repeats",
            12,
            FMT_INT,
            cnt(f",{col('repeat_counted')},1"),
        ),  # fmt: skip
        Column("repeats", "Repeat faults", 9, FMT_INT, cnt(f",{col('is_repeat')},1")),
        Column(
            "rep_rate",
            "Repeat rate",
            9,
            FMT_PCT,
            lambda r, c: f'=IF({c["rep_counted"]}{r}=0,"",{c["repeats"]}{r}/{c["rep_counted"]}{r})',
        ),
        Column(
            "vs_fleet",
            "Repeat rate vs fleet",
            10,
            FMT_X,
            lambda r, c: (
                f'=IF(OR({c["rep_rate"]}{r}="",${c["rep_rate"]}${c["_total"]}=0),"",'
                f"{c['rep_rate']}{r}/${c['rep_rate']}${c['_total']})"
            ),
        ),
        Column("open", "Open tickets", 9, FMT_INT, cnt(f",{col('is_open')},1")),
        Column("over30", "Open > 30 days", 10, FMT_INT, cnt(f',{col("age_days")},">30"')),
    ]
    L = _letters(columns)
    total = {
        "tickets": f"=COUNTA({col('report_id')})",
        "resolved": f"=COUNTIFS({col('resolved')},1)",
        "median": f"=MEDIAN({col('days_to_resolve')})",
        "mean": f"=AVERAGE({col('days_to_resolve')})",
        "breaches": f"=COUNTIFS({col('sla_breach')},1)",
        "breach_rate": _ratio(
            f"{L['breaches']}{first_total}", f'COUNTIFS({col("sla_breach")},">=0")'
        ),
        "rep_counted": f"=COUNTIFS({col('repeat_counted')},1)",
        "repeats": f"=COUNTIFS({col('is_repeat')},1)",
        "rep_rate": _ratio(f"{L['repeats']}{first_total}", f"{L['rep_counted']}{first_total}"),
        "open": f"=COUNTIFS({col('is_open')},1)",
        "over30": f'=COUNTIFS({col("age_days")},">30")',
    }
    return columns, total


def _build_by_vendor(wb: Workbook, m: Metrics, layout: Layout) -> None:
    keys = list(m.mttr["vendor"]["vendor"])
    first_total = str(6 + len(keys))
    columns, total = _vendor_columns(first_total)
    L, index, total_row = _write_summary(
        wb, "By_Vendor", "Vendor scorecard",
        "Median first, mean beside it. Groups under MinSample resolved tickets are marked "
        "'low sample'. SLA targets come from the Vendors sheet (XLOOKUP).",
        columns, keys, total, layout,
    )  # fmt: skip
    ws = wb["By_Vendor"]
    first, last = 6, total_row - 1
    _cf_low_sample(ws, L, first, last)
    _cf_scale(ws, f"{L['median']}{first}:{L['median']}{last}")
    _cf_bars(ws, f"{L['tickets']}{first}:{L['tickets']}{last}")
    rate = f"{L['breach_rate']}{first}"
    _cf_alert(ws, f"{L['breach_rate']}{first}:{L['breach_rate']}{last}",
              f"AND(ISNUMBER({rate}),{rate}>BreachAlert)")  # fmt: skip
    vs = f"{L['vs_fleet']}{first}"
    _cf_alert(ws, f"{L['vs_fleet']}{first}:{L['vs_fleet']}{last}",
              f"AND(ISNUMBER({vs}),{vs}>=RepeatAlert)", "amber")  # fmt: skip


def _build_by_system(wb: Workbook, m: Metrics, layout: Layout) -> None:
    keys = list(m.repeat_by_system["system_id"])
    keys += [s for s in sorted(m.prepared["system_id"].unique()) if s not in keys]
    first_total = str(6 + len(keys))

    def dlookup(field_name: str):
        return lambda r, c: (
            f'=_xlfn.XLOOKUP(${c["system"]}{r},{col("system_id")},{col(field_name)},"")'
        )

    def cnt(extra: str):
        return lambda r, c: f"=COUNTIFS({col('system_id')},${c['system']}{r}{extra})"

    columns = [
        Column("system", "System", 11),
        Column("system_type", "System type", 22, None, dlookup("system_type")),
        Column("vendor", "Vendor", 11, None, dlookup("vendor")),
        Column("site", "Site", 13, None, dlookup("site")),
        Column("tickets", "Tickets", 9, FMT_INT, cnt("")),
        Column("resolved", "Resolved (n)", 10, FMT_INT, cnt(f",{col('resolved')},1")),
        Column(
            "median",
            "Median days to resolve",
            12,
            FMT_1DP,
            lambda r, c: (
                f'=IF({c["resolved"]}{r}=0,"",MEDIAN(IF({col("system_id")}=${c["system"]}{r},'
                f"{col('days_to_resolve')})))"
            ),
            True,
        ),
        Column(
            "mean",
            "Mean days to resolve",
            12,
            FMT_1DP,
            lambda r, c: (
                f"=IFERROR(AVERAGEIFS({col('days_to_resolve')},{col('system_id')},"
                f'${c["system"]}{r},{col("resolved")},1),"")'
            ),
        ),
        Column(
            "sample",
            "Sample",
            11,
            None,
            lambda r, c: f'=IF({c["resolved"]}{r}<MinSample,"low sample","")',
        ),  # fmt: skip
        Column(
            "rep_counted",
            "Faults counted for repeats",
            12,
            FMT_INT,
            cnt(f",{col('repeat_counted')},1"),
        ),  # fmt: skip
        Column("repeats", "Repeat faults", 9, FMT_INT, cnt(f",{col('is_repeat')},1")),
        Column(
            "rep_rate",
            "Repeat rate",
            9,
            FMT_PCT,
            lambda r, c: f'=IF({c["rep_counted"]}{r}=0,"",{c["repeats"]}{r}/{c["rep_counted"]}{r})',
        ),
        Column(
            "fleet",
            "Fleet repeat rate",
            10,
            FMT_PCT,
            lambda r, c: f"=${c['rep_rate']}${c['_total']}",
        ),  # fmt: skip
        Column(
            "vs_pts",
            "Repeat rate vs fleet (pts)",
            11,
            FMT_PTS,
            lambda r, c: f'=IF({c["rep_rate"]}{r}="","",({c["rep_rate"]}{r}-{c["fleet"]}{r})*100)',
        ),
        Column(
            "vs_x",
            "Repeat rate vs fleet (x)",
            11,
            FMT_X,
            lambda r, c: (
                f'=IF(OR({c["rep_rate"]}{r}="",{c["fleet"]}{r}=0),"",{c["rep_rate"]}{r}/{c["fleet"]}{r})'
            ),
        ),
        Column("open", "Open tickets", 9, FMT_INT, cnt(f",{col('is_open')},1")),
        Column("over30", "Open > 30 days", 10, FMT_INT, cnt(f',{col("age_days")},">30"')),
    ]
    L = _letters(columns)
    total = {
        "tickets": f"=COUNTA({col('report_id')})",
        "resolved": f"=COUNTIFS({col('resolved')},1)",
        "median": f"=MEDIAN({col('days_to_resolve')})",
        "mean": f"=AVERAGE({col('days_to_resolve')})",
        "rep_counted": f"=COUNTIFS({col('repeat_counted')},1)",
        "repeats": f"=COUNTIFS({col('is_repeat')},1)",
        "rep_rate": _ratio(f"{L['repeats']}{first_total}", f"{L['rep_counted']}{first_total}"),
        "open": f"=COUNTIFS({col('is_open')},1)",
        "over30": f'=COUNTIFS({col("age_days")},">30")',
    }
    L, index, total_row = _write_summary(
        wb, "By_System", "System view: resolution time and repeat rate vs the fleet",
        "Sorted by repeat rate at build time (low-sample systems last). 'vs fleet' compares a "
        "system's repeat rate with all repeats / all counted faults.",
        columns, keys, total, layout,
    )  # fmt: skip
    ws = wb["By_System"]
    first, last = 6, total_row - 1
    _cf_low_sample(ws, L, first, last)
    _cf_scale(ws, f"{L['rep_rate']}{first}:{L['rep_rate']}{last}")
    _cf_bars(ws, f"{L['tickets']}{first}:{L['tickets']}{last}")
    vs = f"{L['vs_x']}{first}"
    _cf_alert(ws, f"{L['vs_x']}{first}:{L['vs_x']}{last}",
              f"AND(ISNUMBER({vs}),{vs}>=RepeatAlert)", "amber")  # fmt: skip


def _build_by_site(wb: Workbook, m: Metrics, layout: Layout) -> None:
    keys = list(m.mttr["site"]["site"])
    first_total = str(6 + len(keys))

    def cnt(extra: str):
        return lambda r, c: f"=COUNTIFS({col('site')},${c['site']}{r}{extra})"

    columns = [
        Column("site", "Site", 14),
        Column("tickets", "Tickets", 9, FMT_INT, cnt("")),
        Column("resolved", "Resolved (n)", 10, FMT_INT, cnt(f",{col('resolved')},1")),
        Column(
            "median",
            "Median days to resolve",
            12,
            FMT_1DP,
            lambda r, c: (
                f'=IF({c["resolved"]}{r}=0,"",MEDIAN(IF({col("site")}=${c["site"]}{r},'
                f"{col('days_to_resolve')})))"
            ),
            True,
        ),
        Column(
            "mean",
            "Mean days to resolve",
            12,
            FMT_1DP,
            lambda r, c: (
                f"=IFERROR(AVERAGEIFS({col('days_to_resolve')},{col('site')},${c['site']}{r},"
                f'{col("resolved")},1),"")'
            ),
        ),
        Column(
            "sample",
            "Sample",
            11,
            None,
            lambda r, c: f'=IF({c["resolved"]}{r}<MinSample,"low sample","")',
        ),  # fmt: skip
        Column("breaches", "SLA breaches", 10, FMT_INT, cnt(f",{col('sla_breach')},1")),
        Column(
            "breach_rate",
            "SLA breach rate",
            10,
            FMT_PCT,
            lambda r, c: f'=IF({c["resolved"]}{r}=0,"",{c["breaches"]}{r}/{c["resolved"]}{r})',
        ),
        Column("open", "Open tickets", 9, FMT_INT, cnt(f",{col('is_open')},1")),
        Column("over30", "Open > 30 days", 10, FMT_INT, cnt(f',{col("age_days")},">30"')),
        Column(
            "median_age",
            "Median age of open tickets (days)",
            14,
            FMT_1DP,
            lambda r, c: (
                f'=IFERROR(MEDIAN(IF({col("site")}=${c["site"]}{r},{col("age_days")})),"")'
            ),
            True,
        ),
        Column(
            "rep_counted",
            "Faults counted for repeats",
            12,
            FMT_INT,
            cnt(f",{col('repeat_counted')},1"),
        ),  # fmt: skip
        Column("repeats", "Repeat faults", 9, FMT_INT, cnt(f",{col('is_repeat')},1")),
        Column(
            "rep_rate",
            "Repeat rate",
            9,
            FMT_PCT,
            lambda r, c: f'=IF({c["rep_counted"]}{r}=0,"",{c["repeats"]}{r}/{c["rep_counted"]}{r})',
        ),
    ]
    L = _letters(columns)
    total = {
        "tickets": f"=COUNTA({col('report_id')})",
        "resolved": f"=COUNTIFS({col('resolved')},1)",
        "median": f"=MEDIAN({col('days_to_resolve')})",
        "mean": f"=AVERAGE({col('days_to_resolve')})",
        "breaches": f"=COUNTIFS({col('sla_breach')},1)",
        "breach_rate": _ratio(
            f"{L['breaches']}{first_total}", f'COUNTIFS({col("sla_breach")},">=0")'
        ),
        "open": f"=COUNTIFS({col('is_open')},1)",
        "over30": f'=COUNTIFS({col("age_days")},">30")',
        "median_age": f'=IFERROR(MEDIAN({col("age_days")}),"")',
        "rep_counted": f"=COUNTIFS({col('repeat_counted')},1)",
        "repeats": f"=COUNTIFS({col('is_repeat')},1)",
        "rep_rate": _ratio(f"{L['repeats']}{first_total}", f"{L['rep_counted']}{first_total}"),
    }
    L, index, total_row = _write_summary(
        wb, "By_Site", "Site view",
        "Open > 30 days is the backlog that has aged; the median age shows how old the typical "
        "open ticket is.",
        columns, keys, total, layout,
    )  # fmt: skip
    ws = wb["By_Site"]
    first, last = 6, total_row - 1
    _cf_low_sample(ws, L, first, last)
    _cf_scale(ws, f"{L['median']}{first}:{L['median']}{last}")
    _cf_bars(ws, f"{L['over30']}{first}:{L['over30']}{last}")
    rate = f"{L['breach_rate']}{first}"
    _cf_alert(ws, f"{L['breach_rate']}{first}:{L['breach_rate']}{last}",
              f"AND(ISNUMBER({rate}),{rate}>BreachAlert)")  # fmt: skip


def _build_monthly(wb: Workbook, m: Metrics, layout: Layout) -> None:
    months = list(m.monthly["month"])

    def month_end(r, c):
        return f"MIN(EOMONTH(${c['month']}{r},0),AsOfDate)"

    columns = [
        Column(
            "month",
            "Month",
            12,
            FMT_MONTH,
            lambda r, c: (
                f"=DATE(YEAR(MIN({col('date_opened')})),MONTH(MIN({col('date_opened')})),1)"
                if r == 6
                else f"=EDATE({c['month']}{r - 1},1)"
            ),
        ),
        Column(
            "opened",
            "Opened",
            9,
            FMT_INT,
            lambda r, c: f"=COUNTIFS({col('opened_month')},${c['month']}{r})",
        ),  # fmt: skip
        Column(
            "closed",
            "Closed (resolved)",
            10,
            FMT_INT,
            lambda r, c: f"=COUNTIFS({col('closed_month')},${c['month']}{r})",
        ),  # fmt: skip
        Column(
            "net", "Net opened", 9, "+0;-0;0", lambda r, c: f"={c['opened']}{r}-{c['closed']}{r}"
        ),  # fmt: skip
        Column(
            "median",
            "Median days (tickets closed that month)",
            14,
            FMT_1DP,
            lambda r, c: (
                f'=IF({c["closed"]}{r}=0,"",MEDIAN(IF({col("closed_month")}=${c["month"]}{r},'
                f"{col('days_to_resolve')})))"
            ),
            True,
        ),
        Column(
            "mean",
            "Mean days",
            10,
            FMT_1DP,
            lambda r, c: (
                f"=IFERROR(AVERAGEIFS({col('days_to_resolve')},{col('closed_month')},"
                f'${c["month"]}{r}),"")'
            ),
        ),
        Column(
            "sample",
            "Sample",
            11,
            None,
            lambda r, c: f'=IF({c["closed"]}{r}<MinSample,"low sample","")',
        ),  # fmt: skip
        Column(
            "open_end",
            "Open at month end",
            12,
            FMT_INT,
            lambda r, c: (
                f'=COUNTIFS({col("date_opened")},"<="&{month_end(r, c)},{col("is_open")},1)'
                f'+COUNTIFS({col("date_opened")},"<="&{month_end(r, c)},{col("resolved")},1,'
                f'{col("date_closed")},">"&{month_end(r, c)})'
            ),
        ),
    ]
    L, index, _ = _write_summary(
        wb, "Monthly", "Monthly trend",
        "Opened vs closed per month; resolution time of the tickets closed that month; backlog at "
        "each month end (tickets closed with an unknown date are left out).",
        columns, months, None, layout,
    )  # fmt: skip
    ws = wb["Monthly"]
    first, last = 6, 5 + len(months)
    _cf_low_sample(ws, L, first, last)
    _cf_bars(ws, f"{L['opened']}{first}:{L['opened']}{last}")


def _build_backlog(wb: Workbook, m: Metrics, layout: Layout) -> None:
    ws = wb.create_sheet("Backlog")
    _title(
        ws, "Open backlog by age",
        "Open = status Open or In Progress. Age = reporting date (Settings) minus open date.",
    )  # fmt: skip
    headers = [
        "Group",
        "Open tickets",
        *BUCKETS,
        "Open > 30 days",
        "Median age (days)",
        "Oldest (days)",
    ]
    letters = {"group": "A", "open": "B", "b0": "C", "b1": "D", "b2": "E", "b3": "F",
               "over30": "G", "median_age": "H", "oldest": "I"}  # fmt: skip
    layout.cols["Backlog"] = letters
    index: dict[str, int] = {}

    def block(first_header_row: int, label: str, field_name: str | None, groups: list[str]):
        _header_row(ws, first_header_row, [label, *headers[1:]])
        for i, group in enumerate(groups):
            r = first_header_row + 1 + i
            index[f"{field_name or 'all'}:{group}"] = r
            _put(ws, f"A{r}", group, font=f(bold=True))
            crit = f",{col(field_name)},$A{r}" if field_name else ""
            _put(ws, f"B{r}", f"=COUNTIFS({col('is_open')},1{crit})", fmt=FMT_INT)
            for letter in ("C", "D", "E", "F"):
                _put(ws, f"{letter}{r}",
                     f"=COUNTIFS({col('age_bucket')},{letter}${first_header_row}{crit})",
                     fmt=FMT_INT)  # fmt: skip
            _put(ws, f"G{r}", f"=E{r}+F{r}", fmt=FMT_INT)
            cond = f"{col(field_name)}=$A{r}" if field_name else None
            median = (
                f"MEDIAN(IF({cond},{col('age_days')}))" if cond else f"MEDIAN({col('age_days')})"
            )
            oldest = f"MAX(IF({cond},{col('age_days')}))" if cond else f"MAX({col('age_days')})"
            _put(ws, f"H{r}", f'=IFERROR({median},"")', array=bool(cond), fmt=FMT_1DP)
            _put(ws, f"I{r}", f"={oldest}", array=bool(cond), fmt=FMT_INT)
        return first_header_row + 1, first_header_row + len(groups)

    block(5, "Scope", None, ["All open tickets"])
    sites = list(m.mttr["site"]["site"])
    vendors = list(m.mttr["vendor"]["vendor"])
    s_first, s_last = block(9, "Site", "site", sites)
    v_first, v_last = block(10 + len(sites) + 2, "Vendor", "vendor", vendors)
    layout.rows["Backlog"] = {"index": index, "overall": 6}
    for rng in (f"C{s_first}:F{s_last}", f"C{v_first}:F{v_last}"):
        _cf_scale(ws, rng)
    _cf_bars(ws, f"G{s_first}:G{s_last}")
    _cf_bars(ws, f"G{v_first}:G{v_last}")
    _widths(ws, {"A": 18, "B": 11, "C": 11, "D": 11, "E": 11, "F": 11, "G": 12, "H": 13, "I": 11})
    ws.sheet_properties.tabColor = BLUE


# ---------------------------------------------------------------------------------------------
# Data quality (values: an audit trail written by the pipeline)
# ---------------------------------------------------------------------------------------------
def load_dq(directory: Path) -> dict | None:
    summary = directory / "summary.json"
    if not summary.exists():
        return None
    return {
        "summary": json.loads(summary.read_text()),
        "ledger": pd.read_csv(directory / "ledger.csv"),
        "unknowns": pd.read_csv(directory / "unknowns.csv"),
    }


def _build_dq(wb: Workbook, dq: dict | None) -> None:
    ws = wb.create_sheet("Data_Quality")
    _title(ws, "Data quality: what the pipeline fixed",
           "Values written by the cleaning step (an audit log, not formulas).")  # fmt: skip
    ws.sheet_properties.tabColor = MUTED
    if dq is None:
        ws["A4"] = "Run `field-ops clean` first to produce the data-quality tables."
        return
    s = dq["summary"]
    rows = [
        ("Raw rows", s["summary"]["raw_rows"]),
        ("Exact duplicates removed", s["summary"]["exact_duplicates_removed"]),
        ("Near-duplicate tickets merged", s["summary"]["near_duplicates_removed"]),
        ("Clean rows delivered", s["summary"]["final_rows"]),
        ("Rows with a Unknown label after cleaning", s["summary"]["rows_with_any_unknown"]),
        ("Low-confidence fuzzy matches flagged", s["low_confidence"]["rows"]),
        ("Blank categories filled from text", s["category_source"].get("filled_from_text", 0)),
        ("Blank categories left Unclassified", s["category_source"].get("unclassified", 0)),
        ("Logged category contradicts the text (flagged, unchanged)", s["mismatch_flagged"]),
    ]
    for i, (label, value) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=label).font = f(bold=True)
        cell = ws.cell(row=i, column=4, value=int(value))
        cell.font, cell.number_format = f(), FMT_INT
    start = 4 + len(rows) + 2
    ws.cell(row=start - 1, column=1, value="What was fixed").font = f(bold=True, size=11)
    _header_row(ws, start, ["Step", "Rule", "", "Rows", "Example"])
    ledger = dq["ledger"].fillna("")
    for i, rec in enumerate(ledger.itertuples(index=False), start=start + 1):
        ws.cell(row=i, column=1, value=rec.step).font = f()
        ws.cell(row=i, column=2, value=rec.rule).font = f()
        cell = ws.cell(row=i, column=4, value=int(rec.rows))
        cell.font, cell.number_format = f(), FMT_INT
        ws.cell(row=i, column=5, value=str(rec.example)).font = f(color=INK2)
    unk_start = start + len(ledger) + 3
    ws.cell(row=unk_start - 1, column=1, value="Unknown values and why").font = f(
        bold=True, size=11
    )
    _header_row(ws, unk_start, ["Field", "Reason", "", "Rows", "Share of rows"])
    for i, rec in enumerate(dq["unknowns"].itertuples(index=False), start=unk_start + 1):
        ws.cell(row=i, column=1, value=rec.field).font = f()
        ws.cell(row=i, column=2, value=rec.reason).font = f()
        ws.cell(row=i, column=4, value=int(rec.rows)).font = f()
        share = ws.cell(row=i, column=5, value=float(rec.share_of_rows))
        share.font, share.number_format = f(), FMT_PCT
    _widths(ws, {"A": 14, "B": 70, "C": 2, "D": 10, "E": 60})


# ---------------------------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------------------------
def _bar_chart(ws, title: str, y_title: str, anchor: str, cats, data, horizontal=True,
               from_rows=False, color=BLUE, label_fmt: str = "0.0") -> None:  # fmt: skip
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.title = title
    chart.add_data(data, titles_from_data=False, from_rows=from_rows)
    chart.set_categories(cats)
    chart.legend = None
    chart.gapWidth = 60
    series = chart.series[0]
    series.graphicalProperties.solidFill = color
    series.graphicalProperties.line.noFill = True
    chart.dataLabels = DataLabelList()
    labels = chart.dataLabels
    labels.showVal = True
    labels.showSerName = labels.showCatName = labels.showLegendKey = labels.showPercent = False
    labels.numFmt = label_fmt
    for axis in (chart.x_axis, chart.y_axis):
        axis.delete = False
    chart.y_axis.title = y_title
    chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=GRID))
    chart.x_axis.tickLblSkip = 1  # show every category label
    if horizontal:
        chart.x_axis.scaling.orientation = "maxMin"  # largest value at the top
        chart.y_axis.crosses = "max"  # ...and keep the value axis at the bottom
    chart.width, chart.height = 13.2, 9.0 if horizontal else 8.2
    ws.add_chart(chart, anchor)


def _build_dashboard(wb: Workbook, m: Metrics, layout: Layout, selection: tuple[str, str]) -> None:
    ws = wb["Dashboard"]
    ws.sheet_view.showGridLines = False
    ws["B1"] = "Field Operations Analytics"
    ws["B1"].font = f(size=18, bold=True, color=INK)
    ws["B2"] = '="Fault reports | reporting date "&TEXT(AsOfDate,"yyyy-mm-dd")&" | synthetic data"'
    ws["B2"].font = f(color=INK2)

    ws["B4"], ws["E4"] = "Vendor", "Site"
    for ref in ("B4", "E4"):
        ws[ref].font = f(bold=True)
        ws[ref].alignment = Alignment(horizontal="right")
    for ref, value, lst in (("C4", selection[0], "lstVendors"), ("F4", selection[1], "lstSites")):
        ws[ref] = value
        ws[ref].font = f(bold=True, color=INPUT_INK)
        ws[ref].fill = fill(INPUT_FILL)
        ws[ref].border = Border(*(Side(style="thin", color=CARD_LINE),) * 4)
        dv = DataValidation(type="list", formula1=f"={lst}", allow_blank=False)
        dv.error, dv.errorTitle = "Pick a value from the list.", "Not allowed"
        dv.prompt, dv.promptTitle = "Choose from the dropdown.", "Filter"
        ws.add_data_validation(dv)
        dv.add(ref)
    ws["B5"], ws["C5"], ws["F5"] = "filter criteria", '=IF($C$4="All vendors","*",$C$4)', (
        '=IF($F$4="All sites","*",$F$4)')  # fmt: skip
    for ref in ("B5", "C5", "F5"):
        ws[ref].font = f(size=8, color=MUTED)
    ws["B5"].alignment = Alignment(horizontal="right")

    vend, site = "$C$5", "$F$5"
    sel = (f'(({col("vendor")}=$C$4)+($C$4="{ALL_VENDORS}")>0)*'
           f'(({col("site")}=$F$4)+($F$4="{ALL_SITES}")>0)')  # fmt: skip
    base = f"{col('vendor')},{vend},{col('site')},{site}"
    cards = [
        ("tickets", "TICKETS", f"=COUNTIFS({base})", FMT_INT, False, "all reports"),
        ("resolved", "RESOLVED", f"=COUNTIFS({base},{col('resolved')},1)", FMT_INT, False,
         "usable close date"),
        ("open", "OPEN", f"=COUNTIFS({base},{col('is_open')},1)", FMT_INT, False,
         "Open or In Progress"),
        ("open30", "OPEN > 30 DAYS", f'=COUNTIFS({base},{col("age_days")},">30")', FMT_INT,
         False, "aged backlog"),
        ("median", "MEDIAN DAYS TO RESOLVE",
         f'=IFERROR(MEDIAN(IF({sel}*({col("resolved")}=1),{col("days_to_resolve")})),"")',
         FMT_1DP, True, "resolved tickets"),
        ("mean", "MEAN DAYS TO RESOLVE",
         f'=IFERROR(AVERAGEIFS({col("days_to_resolve")},{base},{col("resolved")},1),"")',
         FMT_1DP, False, "pulled up by long tails"),
        ("breach", "SLA BREACH RATE",
         f'=IFERROR(COUNTIFS({base},{col("sla_breach")},1)/COUNTIFS({base},'
         f'{col("sla_breach")},">=0"),"")', FMT_PCT, False, "of resolved tickets"),
        ("repeat", "REPEAT RATE",
         f'=IFERROR(COUNTIFS({base},{col("is_repeat")},1)/COUNTIFS({base},'
         f'{col("repeat_counted")},1),"")', FMT_PCT, False, None),
    ]  # fmt: skip
    for i, (key, label, formula, fmt, array, note) in enumerate(cards):
        letter = get_column_letter(2 + i)
        layout.kpi[key] = f"{letter}8"
        box = Border(*(Side(style="thin", color=CARD_LINE),) * 4)
        for r in (7, 8, 9):
            ws[f"{letter}{r}"].fill = fill(CARD_FILL)
            ws[f"{letter}{r}"].border = box
        _put(ws, f"{letter}7", label, font=f(size=8, bold=True, color=INK2),
             alignment=Alignment(wrap_text=True, vertical="center"))  # fmt: skip
        _put(ws, f"{letter}8", formula, array=array, fmt=fmt,
             font=f(size=22, bold=True, color=INK),
             alignment=Alignment(horizontal="left"))  # fmt: skip
        if note is None:
            note = '="within "&RepeatWindow&" days after close"'
        _put(ws, f"{letter}9", note, font=f(size=8, color=MUTED),
             alignment=Alignment(wrap_text=True, vertical="top"))  # fmt: skip
        ws.column_dimensions[letter].width = 17
    ws.row_dimensions[7].height = 26
    ws.row_dimensions[8].height = 36
    ws.row_dimensions[9].height = 24
    layout.selectors = {"vendor": "C4", "site": "F4"}

    ws["B11"] = (
        '=IF(C8<MinSample,"Low sample: fewer than "&MinSample&" resolved tickets in this '
        'selection, so treat the medians with caution.","")'
    )
    ws["B11"].font = f(bold=True, color=AMBER_INK)
    ws.conditional_formatting.add(
        "H8", FormulaRule(formula=["AND(ISNUMBER($H$8),$H$8>BreachAlert)"], fill=fill(RED_FILL),
                          font=Font(name=FONT, size=22, bold=True, color=RED_INK)),
    )  # fmt: skip
    ws["A1"].font = f()
    ws.column_dimensions["A"].width = 2

    # charts read the summary sheets (whole fleet, not the filters)
    bv = wb["By_Vendor"]
    v = layout.rows["By_Vendor"]["index"]
    vl = layout.cols["By_Vendor"]
    v_first, v_last = min(v.values()), max(v.values())
    _bar_chart(ws, "Median days to resolve, by vendor", "days", "B13",
               Reference(bv, min_col=1, min_row=v_first, max_row=v_last),
               Reference(bv, min_col=_col_idx(vl["median"]), min_row=v_first,
                         max_row=v_last))  # fmt: skip

    mo = wb["Monthly"]
    ml = layout.cols["Monthly"]
    m_first, m_last = 6, 5 + len(m.monthly)
    line = LineChart()
    line.title = "Faults opened vs closed per month"
    for key, color in (("opened", BLUE), ("closed", ORANGE)):
        idx = _col_idx(ml[key])
        line.add_data(Reference(mo, min_col=idx, min_row=m_first - 1, max_row=m_last),
                      titles_from_data=True)  # fmt: skip
        s = line.series[-1]
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = 28575
        s.marker = Marker(symbol="circle", size=7)
        s.marker.graphicalProperties = GraphicalProperties(
            solidFill=color, ln=LineProperties(solidFill=color)
        )
        s.smooth = False
    line.set_categories(Reference(mo, min_col=1, min_row=m_first, max_row=m_last))
    line.x_axis.number_format = "mmm yy"
    for axis in (line.x_axis, line.y_axis):
        axis.delete = False
    line.y_axis.title = "tickets"
    line.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=GRID))
    line.legend.position = "b"
    line.width, line.height = 13.2, 8.2
    ws.add_chart(line, "F13")

    bl = wb["Backlog"]
    _bar_chart(ws, "Open backlog by age (days since opened)", "tickets", "B32",
               Reference(bl, min_col=3, max_col=6, min_row=5),
               Reference(bl, min_col=3, max_col=6, min_row=6, max_row=6),
               horizontal=False, from_rows=True, label_fmt="0")  # fmt: skip

    sy = wb["By_System"]
    sl = layout.cols["By_System"]
    s_first = layout.rows["By_System"]["first"]
    _bar_chart(ws, "Top 10 systems by repeat rate (at build time)", "repeat rate", "F32",
               Reference(sy, min_col=1, min_row=s_first, max_row=s_first + 9),
               Reference(sy, min_col=_col_idx(sl["rep_rate"]), min_row=s_first,
                         max_row=s_first + 9),
               label_fmt="0%")  # fmt: skip
    ws.sheet_properties.tabColor = ORANGE


def _col_idx(letter: str) -> int:
    from openpyxl.utils import column_index_from_string

    return column_index_from_string(letter)


# ---------------------------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------------------------
def build_workbook(
    m: Metrics,
    vendors: pd.DataFrame,
    path: Path,
    dq: dict | None = None,
    selection: tuple[str, str] = (ALL_VENDORS, ALL_SITES),
) -> Layout:
    """Write the report to ``path`` and return where things are (for tests and callers)."""
    wb = Workbook()
    wb.active.title = "Dashboard"
    layout = Layout(path=path)
    _build_settings(wb, m)
    _build_lists(wb, m)
    _build_vendors(wb, vendors)
    _build_data(wb, m, layout)
    _build_by_vendor(wb, m, layout)
    _build_by_system(wb, m, layout)
    _build_by_site(wb, m, layout)
    _build_monthly(wb, m, layout)
    _build_backlog(wb, m, layout)
    _build_dq(wb, dq)
    _build_dashboard(wb, m, layout, selection)

    order = ["Dashboard", "By_Vendor", "By_System", "By_Site", "Monthly", "Backlog", "Vendors",
             "Data", "Data_Quality", "Settings", "Lists"]  # fmt: skip
    wb._sheets = [wb[name] for name in order]
    wb.calculation = CalcProperties(fullCalcOnLoad=True)  # Excel computes everything on open
    wb.properties.title = "Field Operations Analytics report"
    wb.properties.creator = "field-ops-analytics (synthetic data)"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return layout
