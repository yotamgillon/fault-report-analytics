"""Recalculate a workbook with headless LibreOffice and look for formula errors.

openpyxl writes formulas but never calculates them, so an unsupported or mistyped function
would only show up as #NAME? when someone opens the file. Recalculating in LibreOffice (in CI as
well as locally) turns that into a failing check. XLOOKUP needs LibreOffice 24.8 or newer.
"""

import os
import re
import shutil
import subprocess
import tempfile
import warnings
from pathlib import Path

from openpyxl import load_workbook

MIN_VERSION = (24, 8)  # first LibreOffice release that evaluates XLOOKUP
REQUIRE_ENV = "FIELD_OPS_REQUIRE_LIBREOFFICE"
ERROR_VALUES = {"#NAME?", "#VALUE!", "#REF!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!"}
MAC_BINARY = "/Applications/LibreOffice.app/Contents/MacOS/soffice"


class LibreOfficeUnavailable(RuntimeError):
    """LibreOffice is missing or too old for the functions the workbook uses."""


def required() -> bool:
    """CI sets this so a missing LibreOffice fails the build instead of skipping the check."""
    return os.environ.get(REQUIRE_ENV) == "1"


def find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        if path := shutil.which(name):
            return path
    return MAC_BINARY if Path(MAC_BINARY).exists() else None


def version_string(soffice: str) -> str:
    """The full banner, e.g. 'LibreOffice 26.8.1.1 bf6ef2662401...', for logs and test headers."""
    out = subprocess.run([soffice, "--version"], capture_output=True, text=True, timeout=60)
    return out.stdout.strip() or "unknown version"


def version(soffice: str) -> tuple[int, int]:
    out = subprocess.run([soffice, "--version"], capture_output=True, text=True, timeout=60)
    match = re.search(r"(\d+)\.(\d+)", out.stdout)
    if not match:
        raise LibreOfficeUnavailable(f"cannot read LibreOffice version from: {out.stdout!r}")
    return int(match.group(1)), int(match.group(2))


def locate() -> str:
    """Path to a usable soffice, or raise LibreOfficeUnavailable with the reason."""
    soffice = find_soffice()
    if soffice is None:
        raise LibreOfficeUnavailable("LibreOffice (soffice) not found on PATH")
    found = version(soffice)
    if found < MIN_VERSION:
        need = ".".join(map(str, MIN_VERSION))
        raise LibreOfficeUnavailable(
            f"LibreOffice {found[0]}.{found[1]} is too old: XLOOKUP needs {need} or newer"
        )
    return soffice


def recalculate(xlsx: Path, outdir: Path, soffice: str | None = None, timeout: int = 300) -> Path:
    """Open ``xlsx`` in headless LibreOffice, calculate every formula, save a copy in ``outdir``."""
    soffice = soffice or locate()
    outdir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as profile:  # private profile: no lock clashes
        cmd = [
            soffice,
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(outdir),
            str(xlsx),
        ]
        run = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
    result = outdir / xlsx.name
    if run.returncode != 0 or not result.exists():
        raise RuntimeError(f"LibreOffice conversion failed: {run.stdout} {run.stderr}")
    return result


def find_errors(recalculated: Path) -> list[tuple[str, str, str]]:
    """Every cell of a recalculated workbook holding an error value: (sheet, cell, value)."""
    found = []
    with warnings.catch_warnings():  # LibreOffice writes extensions openpyxl does not read
        warnings.simplefilter("ignore", UserWarning)
        workbook = load_workbook(recalculated, data_only=True)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and (value in ERROR_VALUES or value.startswith("Err:")):
                    found.append((sheet.title, cell.coordinate, value))
    return found
