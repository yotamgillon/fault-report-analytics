"""The Excel report: structure, real formulas, and Excel-vs-pandas agreement.

The agreement tests recalculate the workbook in headless LibreOffice (CI installs it and sets
FIELD_OPS_REQUIRE_LIBREOFFICE=1) and compare every summary number with the pandas metrics.
"""

import warnings
import zipfile
from datetime import date, datetime

import helpers
import numpy as np
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula

from field_ops import config, excel_report, libreoffice, metrics, quality

BUCKETS = excel_report.BUCKETS


@pytest.fixture(scope="session")
def built(pipeline_run, tmp_path_factory):
    """Workbook built from the seeded pipeline run, plus the metrics it was built from."""
    base = tmp_path_factory.mktemp("excel")
    final = pipeline_run.classified.final
    vendors = metrics.load_vendors()
    m = metrics.compute_all(final, config.AS_OF, config.REPEAT_WINDOW_DAYS, vendors)
    quality.write_tables(pipeline_run.report, base / "dq")
    path = base / "report.xlsx"
    layout = excel_report.build_workbook(m, vendors, path, excel_report.load_dq(base / "dq"))
    return path, layout, m, final, vendors, base


@pytest.fixture(scope="session")
def formulas(built):
    """The workbook as written (formulas, not values)."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return load_workbook(built[0])


def _recalc(path, soffice, base, name):
    out = libreoffice.recalculate(path, base / name, soffice)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return load_workbook(out, data_only=True), out


@pytest.fixture(scope="session")
def recalculated(built, soffice):
    wb, out = _recalc(built[0], soffice, built[5], "recalc_all")
    return wb, out


# --- structure (no LibreOffice needed) ---------------------------------------------------------
def test_expected_sheets_in_order(formulas):
    assert formulas.sheetnames == [
        "Dashboard", "By_Vendor", "By_System", "By_Site", "Monthly", "Backlog", "Vendors",
        "Data", "Data_Quality", "Settings", "Lists",
    ]  # fmt: skip
    assert formulas["Lists"].sheet_state == "hidden"


def test_data_and_vendor_tables(formulas, built):
    layout = built[1]
    assert set(formulas["Data"].tables) == {"tblFaults"}
    assert set(formulas["Vendors"].tables) == {"tblVendors"}
    ref = formulas["Data"].tables["tblFaults"].ref
    assert ref.endswith(str(layout.n_rows + 1)) and ref.startswith("A1:")
    assert layout.n_rows == len(built[3])


def test_named_inputs_exist(formulas):
    names = set(formulas.defined_names)
    assert {"AsOfDate", "MinSample", "BreachAlert", "RepeatAlert", "RepeatWindow"} <= names
    assert {"lstVendors", "lstSites", "lstStatus", "lstSeverity", "lstCategory"} <= names


def _is_formula(value) -> bool:
    return isinstance(value, ArrayFormula) or (isinstance(value, str) and value.startswith("="))


def _text(value) -> str:
    return value.text if isinstance(value, ArrayFormula) else str(value)


def test_every_summary_number_is_a_formula(formulas, built):
    """Nothing a reader sees on a summary sheet is a pasted value (labels and sort order aside)."""
    layout = built[1]
    for sheet in ("By_Vendor", "By_System", "By_Site"):
        ws, letters = formulas[sheet], layout.cols[sheet]
        info = layout.rows[sheet]
        key_col = next(iter(letters.values()))
        for r in info["index"].values():
            for key, letter in letters.items():
                if letter != key_col:
                    cell = ws[f"{letter}{r}"]
                    assert _is_formula(cell.value), f"{sheet}!{letter}{r} ({key}) is a value"
        for key, letter in letters.items():  # the total row may leave meaningless cells blank
            value = ws[f"{letter}{info['total']}"].value
            assert value is None or letter == key_col or _is_formula(value), f"{sheet} total {key}"
    ws, letters = formulas["Monthly"], layout.cols["Monthly"]
    for r in layout.rows["Monthly"]["index"].values():
        for letter in letters.values():
            assert _is_formula(ws[f"{letter}{r}"].value), f"Monthly!{letter}{r}"
    ws = formulas["Backlog"]
    for r in layout.rows["Backlog"]["index"].values():
        for letter in "BCDEFGHI":
            assert _is_formula(ws[f"{letter}{r}"].value), f"Backlog!{letter}{r}"
    for key, cell in layout.kpi.items():
        assert _is_formula(formulas["Dashboard"][cell].value), key


def test_data_helper_columns_are_formulas_in_every_row(formulas, built):
    layout = built[1]
    ws, letters = formulas["Data"], layout.cols["Data"]
    helpers = [k for k, _, _ in excel_report.FORMULA_COLUMNS]
    for key in helpers:
        for r in range(2, layout.n_rows + 2):
            assert _is_formula(ws[f"{letters[key]}{r}"].value), f"Data {key} row {r}"


def test_xlookup_joins_vendor_info(formulas, built):
    letters = built[1].cols["Data"]
    tier = _text(formulas["Data"][f"{letters['vendor_tier']}2"].value)
    assert tier.startswith("=_xlfn.XLOOKUP(") and "tblVendors[tier]" in tier
    sla = _text(formulas["Data"][f"{letters['sla_target_days']}2"].value)
    assert "tblVendors[sla_target_days]" in sla
    by_vendor = built[1].cols["By_Vendor"]
    first = built[1].rows["By_Vendor"]["first"]
    assert "_xlfn.XLOOKUP" in _text(formulas["By_Vendor"][f"{by_vendor['tier']}{first}"].value)


def test_conditional_medians_use_array_formulas_not_aggregate(formulas, built):
    layout = built[1]
    first = layout.rows["By_Vendor"]["first"]
    median = formulas["By_Vendor"][f"{layout.cols['By_Vendor']['median']}{first}"].value
    assert isinstance(median, ArrayFormula) and "MEDIAN(IF(" in median.text
    every_formula = [
        _text(c.value) for ws in formulas.worksheets for row in ws.iter_rows()
        for c in row if _is_formula(c.value)
    ]  # fmt: skip
    assert not any("AGGREGATE" in text for text in every_formula)  # #VALUE! in LibreOffice
    assert sum("MEDIAN(IF(" in text for text in every_formula) > 40
    kinds = {"COUNTIFS", "AVERAGEIFS", "XLOOKUP", "SUMPRODUCT"} - {"SUMPRODUCT"}
    for kind in kinds:
        assert any(kind in text for text in every_formula), kind


def test_dropdowns(formulas, built):
    dash = formulas["Dashboard"].data_validations.dataValidation
    lists = {str(dv.sqref): dv.formula1 for dv in dash}
    assert lists == {"C4": "=lstVendors", "F4": "=lstSites"}
    data = formulas["Data"].data_validations.dataValidation
    assert {dv.formula1 for dv in data} == {"=lstStatus", "=lstSeverity", "=lstCategory"}


def test_conditional_formatting_of_each_kind(formulas):
    kinds = set()
    for name in ("By_Vendor", "By_System", "By_Site", "Backlog", "Data", "Dashboard"):
        for cf in formulas[name].conditional_formatting:
            kinds |= {rule.type for rule in cf.rules}
    assert {"colorScale", "dataBar", "expression"} <= kinds


def test_dashboard_has_four_charts_and_two_selectors(built):
    with zipfile.ZipFile(built[0]) as z:
        charts = [n for n in z.namelist() if n.startswith("xl/charts/chart")]
    assert len(charts) == 4
    assert built[1].selectors == {"vendor": "C4", "site": "F4"}


def test_settings_document_the_assumptions(formulas):
    text = " ".join(
        str(c.value) for row in formulas["Settings"].iter_rows() for c in row if c.value
    )
    for needle in ("synthetic", "Resolved", "Repeat fault", "yellow", "Low sample"):
        assert needle.lower() in text.lower()


def test_font_is_arial_on_summary_sheets(formulas):
    assert formulas["By_Vendor"]["A6"].font.name == "Arial"
    assert formulas["Dashboard"]["B8"].font.name == "Arial"


# --- Excel vs pandas (needs LibreOffice) -------------------------------------------------------
def _close(got, want) -> bool:
    missing = lambda v: v is None or v == "" or (isinstance(v, float) and np.isnan(v))  # noqa: E731
    if missing(got) or missing(want):
        return missing(got) and missing(want)
    if isinstance(got, str) or isinstance(want, str):
        return got == want
    return abs(float(got) - float(want)) < 1e-9


def compare(wb, layout, m: metrics.Metrics) -> list[str]:
    """Every summary cell in the recalculated workbook against the pandas metrics."""
    bad: list[str] = []

    def check(sheet, ref, want, label):
        got = wb[sheet][ref].value
        if not _close(got, want):
            bad.append(f"{sheet}!{ref} {label}: excel={got!r} pandas={want!r}")

    sc = m.scorecard.set_index("vendor")
    letters = layout.cols["By_Vendor"]
    for vendor, r in layout.rows["By_Vendor"]["index"].items():
        want = {"tickets": "tickets", "resolved": "resolved_n", "median": "median_days",
                "mean": "mean_days", "breaches": "sla_breaches", "breach_rate": "sla_breach_rate",
                "rep_counted": "counted_faults", "repeats": "repeats", "rep_rate": "repeat_rate",
                "vs_fleet": "vs_fleet_x", "open": "open", "over30": "over_30_days"}  # fmt: skip
        for key, column in want.items():
            check("By_Vendor", f"{letters[key]}{r}", sc.at[vendor, column], f"{vendor} {key}")
        check("By_Vendor", f"{letters['tier']}{r}", sc.at[vendor, "tier"], f"{vendor} tier")
        check("By_Vendor", f"{letters['sla_days']}{r}", sc.at[vendor, "sla_target_days"], "sla")

    rs = m.repeat_by_system.set_index("system_id")
    letters = layout.cols["By_System"]
    for system, r in layout.rows["By_System"]["index"].items():
        want = {"tickets": "tickets", "resolved": "resolved_n", "median": "median_days",
                "rep_counted": "counted_faults", "repeats": "repeats", "rep_rate": "repeat_rate",
                "vs_pts": "vs_fleet_pts", "vs_x": "vs_fleet_x", "vendor": "vendor",
                "site": "site"}  # fmt: skip
        for key, column in want.items():
            check("By_System", f"{letters[key]}{r}", rs.at[system, column], f"{system} {key}")

    mttr, aging = m.mttr["site"].set_index("site"), m.aging_by_site.set_index("site")
    letters = layout.cols["By_Site"]
    for site, r in layout.rows["By_Site"]["index"].items():
        pairs = {"tickets": mttr.at[site, "tickets"], "resolved": mttr.at[site, "n"],
                 "median": mttr.at[site, "median_days"], "mean": mttr.at[site, "mean_days"],
                 "open": aging.at[site, "open"], "over30": aging.at[site, "over_30_days"],
                 "median_age": aging.at[site, "median_age_days"]}  # fmt: skip
        for key, want in pairs.items():
            check("By_Site", f"{letters[key]}{r}", want, f"{site} {key}")

    monthly = m.monthly.set_index("month")
    letters = layout.cols["Monthly"]
    for month, r in layout.rows["Monthly"]["index"].items():
        if month not in monthly.index:
            continue
        pairs = {
            "opened": "opened",
            "closed": "closed",
            "net": "net_opened",
            "median": "median_days",
            "mean": "mean_days",
            "open_end": "open_at_month_end",
        }
        for key, column in pairs.items():
            check("Monthly", f"{letters[key]}{r}", monthly.at[month, column], f"{month} {key}")
        label = wb["Monthly"][f"{letters['month']}{r}"].value
        if str(label)[:7] != month:
            bad.append(f"Monthly month label {label!r} != {month}")

    overall = m.aging_overall.iloc[0]
    for bucket, letter in zip(BUCKETS, "CDEF", strict=True):
        check("Backlog", f"{letter}6", overall[bucket], f"overall {bucket}")
    for key, r in layout.rows["Backlog"]["index"].items():
        kind, group = key.split(":", 1)
        if kind == "all":
            continue
        table = (m.aging_by_site.set_index("site") if kind == "site"
                 else m.aging_by_vendor.set_index("vendor"))  # fmt: skip
        for bucket, letter in zip(BUCKETS, "CDEF", strict=True):
            check("Backlog", f"{letter}{r}", table.at[group, bucket], f"{key} {bucket}")
        check("Backlog", f"H{r}", table.at[group, "median_age_days"], f"{key} median age")
        check("Backlog", f"I{r}", table.at[group, "oldest_days"], f"{key} oldest")
    return bad


def expected_kpis(m: metrics.Metrics, vendor: str | None, site: str | None) -> dict:
    return helpers.expected_kpis(m, [vendor] if vendor else None, [site] if site else None)


def test_recalculation_has_no_formula_errors(recalculated):
    assert libreoffice.find_errors(recalculated[1]) == []


def test_error_check_can_actually_fail(soffice, tmp_path):
    wb = Workbook()
    wb.active["A1"] = "=1/0"
    wb.active["A2"] = '=_xlfn.XLOOKUP("x",B1:B2,C1:C2)'
    wb.active["A3"] = "=NOT_A_FUNCTION(1)"
    wb.save(tmp_path / "bad.xlsx")
    out = libreoffice.recalculate(tmp_path / "bad.xlsx", tmp_path / "out", soffice)
    found = {cell: value for _, cell, value in libreoffice.find_errors(out)}
    assert found["A1"] == "#DIV/0!"
    assert "A3" in found  # an unknown function surfaces as an error, not a silent blank


def test_xlookup_is_supported_by_this_libreoffice(soffice, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["A2"], ws["B2"] = "a", 1, "b", 2
    ws["D1"] = '=_xlfn.XLOOKUP("b",A1:A2,B1:B2,"none")'
    wb.save(tmp_path / "x.xlsx")
    out = libreoffice.recalculate(tmp_path / "x.xlsx", tmp_path / "o", soffice)
    assert load_workbook(out, data_only=True).active["D1"].value == 2


def test_old_libreoffice_is_rejected(monkeypatch):
    monkeypatch.setattr(libreoffice, "find_soffice", lambda: "/fake/soffice")
    monkeypatch.setattr(libreoffice, "version", lambda _: (24, 2))
    with pytest.raises(libreoffice.LibreOfficeUnavailable, match="XLOOKUP needs 24.8"):
        libreoffice.locate()


def test_summary_sheets_agree_with_pandas(recalculated, built):
    assert compare(recalculated[0], built[1], built[2]) == []


def test_dashboard_kpis_for_the_whole_fleet(recalculated, built):
    wb, layout, m = recalculated[0], built[1], built[2]
    want = expected_kpis(m, None, None)
    for key, cell in layout.kpi.items():
        assert _close(wb["Dashboard"][cell].value, want[key]), key
    assert wb["Dashboard"]["B11"].value in (None, "")  # no low-sample warning for the fleet


def test_data_helper_columns_agree_with_pandas(recalculated, built):
    wb, layout, m = recalculated[0], built[1], built[2]
    ws, letters, p = wb["Data"], layout.cols["Data"], m.prepared
    rows = list(ws.iter_rows(min_row=2, max_row=layout.n_rows + 1, values_only=True))
    index = {key: i for i, key in enumerate(letters)}

    def column(key):
        return [row[index[key]] for row in rows]

    assert column("resolved") == p["is_resolved"].astype(int).tolist()
    assert all(_close(a, b) for a, b in zip(column("days_to_resolve"), p["days_to_resolve"],
                                            strict=True))  # fmt: skip
    assert column("is_open") == p["is_open"].astype(int).tolist()
    assert all(_close(a, b) for a, b in zip(column("age_days"), p["age_days"], strict=True))
    assert all(_close(a, b) for a, b in zip(column("sla_target_days"), p["sla_target_days"],
                                            strict=True))  # fmt: skip
    breach = [None if v == "" else v for v in column("sla_breach")]
    want = [
        None if not r else int(b) for r, b in zip(p["is_resolved"], p["sla_breach"], strict=True)
    ]
    assert breach == want
    tier = dict(zip(built[4]["vendor"], built[4]["tier"], strict=True))
    assert column("vendor_tier") == [tier[v] for v in p["vendor"]]
    buckets = [None if v == "" else v for v in column("age_bucket")]
    assert set(filter(None, buckets)) <= set(BUCKETS)
    assert sum(b is not None for b in buckets) == int(p["is_open"].sum())


def test_dashboard_filters_vendor_and_site(built, soffice):
    path, layout, m, final, vendors, base = built
    selection = ("Vendor F", "Ridgeview")
    selected = base / "selected.xlsx"
    excel_report.build_workbook(m, vendors, selected, selection=selection)
    wb, _ = _recalc(selected, soffice, base, "recalc_selected")
    want = expected_kpis(m, *selection)
    assert want["tickets"] > 0
    for key, cell in layout.kpi.items():
        assert _close(wb["Dashboard"][cell].value, want[key]), (key, wb["Dashboard"][cell].value)
    vendor_only = base / "vendor_only.xlsx"
    excel_report.build_workbook(m, vendors, vendor_only, selection=("Vendor F", "All sites"))
    wb2, _ = _recalc(vendor_only, soffice, base, "recalc_vendor_only")
    want2 = expected_kpis(m, "Vendor F", None)
    for key, cell in layout.kpi.items():
        assert _close(wb2["Dashboard"][cell].value, want2[key]), key


def test_dashboard_warns_about_a_low_sample_selection(built, soffice):
    """A vendor/site pair with almost no tickets: warn, show blanks, and never show an error."""
    path, layout, m, final, vendors, base = built
    p = m.prepared
    counts = p[p["is_resolved"]].groupby(["vendor", "site"]).size().unstack(fill_value=0)
    vendor, site = counts.stack()[lambda c: c < metrics.MIN_SAMPLE].index[0]
    small = base / "small.xlsx"
    excel_report.build_workbook(m, vendors, small, selection=(vendor, site))
    wb, out = _recalc(small, soffice, base, "recalc_small")
    assert libreoffice.find_errors(out) == []
    message = wb["Dashboard"]["B11"].value
    assert message and "Low sample" in message
    want = expected_kpis(m, vendor, site)
    for key, cell in layout.kpi.items():
        assert _close(wb["Dashboard"][cell].value, want[key]), key


def _edit_settings(path, out, **cells):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(path)
    for ref, value in cells.items():
        wb["Settings"][ref] = value
    wb.save(out)


def test_changing_min_sample_flows_through_the_formulas(built, soffice):
    path, layout, m, final, vendors, base = built
    edited = base / "min_sample.xlsx"
    _edit_settings(path, edited, C5=200)  # Settings!C5 is MinSample
    wb, _ = _recalc(edited, soffice, base, "recalc_min_sample")
    letters, rows = layout.cols["By_Vendor"], layout.rows["By_Vendor"]["index"]
    resolved = m.scorecard.set_index("vendor")["resolved_n"]
    for vendor, r in rows.items():
        flagged = wb["By_Vendor"][f"{letters['sample']}{r}"].value == "low sample"
        assert flagged == (resolved[vendor] < 200), vendor
    assert sum(wb["By_Vendor"][f"{letters['sample']}{r}"].value == "low sample"
               for r in rows.values()) > 0  # fmt: skip


def test_changing_the_reporting_date_ages_the_backlog(built, soffice):
    path, layout, m, final, vendors, base = built
    later = date(2026, 10, 31)
    edited = base / "as_of.xlsx"
    _edit_settings(path, edited, C4=datetime(2026, 10, 31))  # Settings!C4 is AsOfDate
    wb, out = _recalc(edited, soffice, base, "recalc_as_of")
    m_later = metrics.compute_all(final, later, config.REPEAT_WINDOW_DAYS, vendors)
    assert libreoffice.find_errors(out) == []
    assert compare(wb, layout, m_later) == []
    assert _close(wb["Dashboard"][layout.kpi["open30"]].value, m_later.kpis["open_over_30_days"])
    assert m_later.kpis["open_over_30_days"] > m.kpis["open_over_30_days"]
